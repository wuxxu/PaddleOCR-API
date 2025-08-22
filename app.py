import os
from flask import Flask, request, send_file
from pdf2image import convert_from_path
from paddleocr import PaddleOCR
from docx import Document

app = Flask(__name__)
ocr = PaddleOCR(use_angle_cls=True, lang='en')

def sort_ocr_results(results):
    # Sort lines by Y coordinate
    lines = []
    for line in results[0]:
        box, (text, conf) = line
        y_min = min([point[1] for point in box])
        lines.append((y_min, text))
    return sorted(lines, key=lambda x: x[0])

@app.route('/pdf-to-docx', methods=['POST'])
def pdf_to_docx():
    pdf_file = request.files['file']
    input_path = "/tmp/input.pdf"
    output_path = "/tmp/output.docx"
    pdf_file.save(input_path)

    images = convert_from_path(input_path)
    doc = Document()

    for page_num, image in enumerate(images, start=1):
        image_path = f"/tmp/page_{page_num}.png"
        image.save(image_path, "PNG")

        results = ocr.ocr(image_path, cls=True)
        sorted_lines = sort_ocr_results(results)

        doc.add_heading(f"Page {page_num}", level=2)

        last_y = None
        paragraph = ""
        for y, text in sorted_lines:
            if last_y is None:
                last_y = y
            # If the line is far below the last one → new paragraph
            if abs(y - last_y) > 20:
                if paragraph:
                    doc.add_paragraph(paragraph)
                paragraph = text
            else:
                paragraph += " " + text
            last_y = y
        if paragraph:
            doc.add_paragraph(paragraph)

        doc.add_page_break()

    doc.save(output_path)
    return send_file(output_path, as_attachment=True, download_name="converted.docx")

@app.route('/')
def home():
    return "OCR PDF → DOCX API (layout-aware) is running 🚀"
